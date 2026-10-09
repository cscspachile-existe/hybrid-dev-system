"""Genera el DOCX de cotización oficial a partir de CotizacionData."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .excel_io import CotizacionData
from .formato_clp import formatear_clp

ROOT = Path(__file__).resolve().parents[1]
LOGO_PATH = ROOT / "plantillas" / "assets" / "logo.png"

# Paleta sobria institucional (Mercado Público / organismo público)
COLOR_TITULO = RGBColor(0x1A, 0x36, 0x5D)
COLOR_TEXTO = RGBColor(0x1F, 0x29, 0x37)
COLOR_TABLA_HEADER_BG = "1A365D"
COLOR_TABLA_HEADER_FG = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_FILA_ALT = "F3F4F6"


def _set_run_font(run, size_pt: float = 11, bold: bool = False, color=COLOR_TEXTO):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = color


def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 0:
        _set_run_font(run, 22, bold=True, color=COLOR_TITULO)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        _set_run_font(run, 12, bold=True, color=COLOR_TITULO)
    else:
        _set_run_font(run, 14, bold=True, color=COLOR_TITULO)
    p.paragraph_format.space_before = Pt(12 if level else 0)
    p.paragraph_format.space_after = Pt(6)
    return p


def _add_body(doc: Document, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        _set_run_font(r1, 11, bold=True)
        r2 = p.add_run(text)
        _set_run_font(r2, 11)
    else:
        r = p.add_run(text)
        _set_run_font(r, 11)
    p.paragraph_format.space_after = Pt(4)
    return p


def _shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = tcPr.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): hex_color,
        },
    )
    tcPr.append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, color=COLOR_TEXTO, center: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    _set_run_font(run, 10, bold=bold, color=color)


def _build_emisor_table(doc: Document, data: CotizacionData):
    filas = [
        ("Identificación del Proveedor", "Detalle Institucional"),
        ("Razón Social", data.emisor.get("razon_social", "")),
        ("Nombre de Fantasía", data.emisor.get("nombre_fantasia", "")),
        ("R.U.T.", data.emisor.get("rut", "")),
        ("Domicilio Legal", data.emisor.get("domicilio_legal", "")),
        ("Condición Proveedor", data.emisor.get("condicion_proveedor", "")),
        ("Contacto Comercial", data.emisor.get("contacto_comercial", "")),
    ]
    table = doc.add_table(rows=len(filas), cols=2)
    table.style = "Table Grid"
    for i, (a, b) in enumerate(filas):
        _set_cell_text(table.rows[i].cells[0], a, bold=True, center=(i == 0))
        _set_cell_text(
            table.rows[i].cells[1],
            b,
            bold=(i == 0),
            color=COLOR_TABLA_HEADER_FG if i == 0 else COLOR_TEXTO,
            center=(i == 0),
        )
        if i == 0:
            _shade_cell(table.rows[i].cells[0], COLOR_TABLA_HEADER_BG)
            _shade_cell(table.rows[i].cells[1], COLOR_TABLA_HEADER_BG)
            _set_cell_text(
                table.rows[i].cells[0], a, bold=True, color=COLOR_TABLA_HEADER_FG, center=True
            )
            _set_cell_text(
                table.rows[i].cells[1], b, bold=True, color=COLOR_TABLA_HEADER_FG, center=True
            )
        elif i % 2 == 0:
            _shade_cell(table.rows[i].cells[0], COLOR_FILA_ALT)
            _shade_cell(table.rows[i].cells[1], COLOR_FILA_ALT)
    return table


def _build_items_table(doc: Document, data: CotizacionData):
    table = doc.add_table(rows=1 + len(data.items), cols=5)
    table.style = "Table Grid"
    headers = [
        "Ítem",
        "Descripción del Producto",
        "Cantidad",
        "Precio Unitario Neto",
        "Subtotal Neto",
    ]
    for i, h in enumerate(headers):
        _set_cell_text(
            table.rows[0].cells[i], h, bold=True, color=COLOR_TABLA_HEADER_FG, center=True
        )
        _shade_cell(table.rows[0].cells[i], COLOR_TABLA_HEADER_BG)

    for r, item in enumerate(data.items, start=1):
        vals = [
            str(item.item or r),
            item.descripcion,
            str(int(item.cantidad) if float(item.cantidad).is_integer() else item.cantidad),
            formatear_clp(item.precio_unitario_neto),
            formatear_clp(item.subtotal_neto),
        ]
        for c, val in enumerate(vals):
            center = c != 1
            _set_cell_text(table.rows[r].cells[c], val, center=center)
            if r % 2 == 0:
                _shade_cell(table.rows[r].cells[c], COLOR_FILA_ALT)
    return table


def generar_docx(
    data: CotizacionData,
    salida: str | Path,
    logo: str | Path | None = None,
) -> Path:
    """Escribe el DOCX de cotización oficial."""
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    logo_path = Path(logo) if logo else LOGO_PATH
    if logo_path.exists():
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = hp.add_run()
        # ancho ~ 8 cm (logo horizontal CSC)
        run.add_picture(str(logo_path), width=Cm(8))

    _add_heading(doc, "COTIZACIÓN OFICIAL", level=0)
    _add_heading(
        doc,
        data.cotizacion.get("titulo_oferta")
        or (
            "Oferta Técnico-Económica para Mercado Público | "
            f"Licitación ID {data.cotizacion.get('id_proceso', '')}"
        ),
        level=2,
    )

    _add_heading(doc, "1. Encabezado y Datos del Emisor", level=1)
    _build_emisor_table(doc, data)

    _add_heading(doc, "2. Antecedentes del Cliente y Proceso", level=1)
    _add_body(doc, data.cotizacion.get("cliente", ""), bold_prefix="Cliente: ")
    _add_body(
        doc,
        data.cotizacion.get("unidad_solicitante", ""),
        bold_prefix="Unidad Solicitante: ",
    )
    _add_body(
        doc,
        data.cotizacion.get("responsable_proceso", "") or "—",
        bold_prefix="Responsable del Proceso: ",
    )
    _add_body(
        doc,
        data.cotizacion.get("referencia_proceso")
        or f"Solicitud de Cotización ID {data.cotizacion.get('id_proceso', '')}",
        bold_prefix="Referencia del Proceso: ",
    )

    _add_heading(doc, "3. Propuesta Económica Detallada", level=1)
    _build_items_table(doc, data)

    _add_heading(doc, "Resumen Financiero", level=1)
    _add_body(doc, formatear_clp(data.subtotal_neto, con_sufijo=False), bold_prefix="Subtotal Neto: ")
    _add_body(
        doc,
        formatear_clp(data.iva),
        bold_prefix=f"IVA ({int(data.iva_pct) if data.iva_pct.is_integer() else data.iva_pct}%): ",
    )
    total_p = _add_body(
        doc,
        formatear_clp(data.total_bruto),
        bold_prefix="TOTAL OFERTADO (Bruto): ",
    )
    for run in total_p.runs:
        run.bold = True

    _add_heading(doc, "4. Matriz de Cumplimiento / Justificación Técnica", level=1)
    just = data.cotizacion.get("justificacion_tecnica", "").strip()
    _add_body(
        doc,
        just
        or (
            "La propuesta cumple con las especificaciones técnicas solicitadas "
            "en el proceso de cotización de Mercado Público."
        ),
    )

    _add_heading(doc, "5. Condiciones Comerciales y de Servicio", level=1)
    _add_body(
        doc,
        data.condiciones.get("plazo_entrega", ""),
        bold_prefix="Plazo de Entrega: ",
    )
    _add_body(
        doc,
        data.condiciones.get("garantia_tecnica", ""),
        bold_prefix="Garantía Técnica: ",
    )
    _add_body(
        doc,
        data.condiciones.get("validez_oferta", ""),
        bold_prefix="Validez de la Oferta: ",
    )
    _add_body(
        doc,
        data.condiciones.get("forma_pago", ""),
        bold_prefix="Forma de Pago: ",
    )

    doc.save(salida)
    return salida
